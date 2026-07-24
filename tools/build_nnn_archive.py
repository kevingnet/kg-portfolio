#!/usr/bin/env python3
"""Import /all/nnn archive into portfolio assets and JSON manifest."""

from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from urllib.parse import unquote

ROOT = Path(__file__).resolve().parent.parent
NNN_ROOT = Path("/media/kg/fecd6373-9e9f-486b-b9b8-f798dc71fc77/all/nnn")
ASSETS_ARCHIVE = ROOT / "assets" / "archive"
MANIFEST = ROOT / "data" / "nnn-archive.json"

NNN_TO_SLUG = {
    "DirecTV": "directv",
    "Disney": "disney",
    "Electrosonic": "electrosonic",
    "Google": "google",
    "HMS": "hms",
    "JakeKnows": "jakeknows",
    "Motorola": "motorola",
    "OpenTV": "opentv",
    "PosDev": "posdev",
    "Spirent": "spirent",
    "Surfware": "surfware",
    "TelVista": "audiotelco",
    "Vmware": "vmware",
    "Volt": "voltdelta",
    "Yahoo": "yahoo",
}

IMG_EXT = {".bmp", ".gif", ".jpg", ".jpeg", ".png", ".pcx", ".tif", ".tiff", ".ico", ".webp"}
CODE_EXT = {
    ".cpp", ".c", ".h", ".hpp", ".cs", ".java", ".py", ".js", ".asp", ".sql", ".php", ".tcl",
}
DOC_EXT = {".doc", ".docx", ".pdf", ".ppt", ".pptx", ".rtf", ".dotx"}
TEXT_EXT = {".txt", ".md", ".arch", ".log", ".htm", ".html", ".srt", ".yaml", ".yml"}
OFFICE_HTML_EXT = {".doc", ".docx", ".dotx", ".rtf"}

LANG_MAP = {
    ".cpp": "cpp", ".c": "c", ".h": "cpp", ".hpp": "cpp", ".cs": "csharp",
    ".java": "java", ".py": "python", ".js": "javascript", ".asp": "html",
    ".sql": "sql", ".php": "php", ".tcl": "tcl", ".htm": "html", ".html": "html",
    ".yaml": "yaml", ".yml": "yaml",
}

MAX_CODE_LINES = 220
MAX_TEXT_CHARS = 14_000
MAX_DOC_CHARS = 18_000
MAX_DOC_SEGMENTS = 120
MAX_DOC_IMAGES = 48
MIN_IMG_DIM = 80
MIN_IMG_BYTES = 1200

WHITESPACE_RE = re.compile(r"\n{3,}")
SAFE_NAME_RE = re.compile(r"[^A-Za-z0-9._-]+")


def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = WHITESPACE_RE.sub("\n\n", text)
    return text.strip()


def truncate(text: str, limit: int) -> tuple[str, bool]:
    if len(text) <= limit:
        return text, False
    return text[: limit - 40].rstrip() + "\n\n… [truncated for display]", True


def truncate_lines(text: str, max_lines: int) -> tuple[str, bool, int]:
    lines = text.splitlines()
    total = len(lines)
    if total <= max_lines:
        return text, False, total
    shown = "\n".join(lines[:max_lines])
    return shown + f"\n\n… [truncated — {total} lines total]", True, total


def classify(path: Path) -> str:
    ext = path.suffix.lower()
    name = path.name.lower()
    if ext in IMG_EXT:
        return "image"
    if ext in CODE_EXT:
        return "code"
    if ext in DOC_EXT or name.endswith(",pdf") or name.endswith(".pdf"):
        return "document"
    if ext in TEXT_EXT or name in ("makefile", "readme") or name.startswith("readme"):
        return "text"
    return "other"


def read_text_file(path: Path) -> str:
    for enc in ("utf-8", "latin-1", "cp1252"):
        try:
            return path.read_text(encoding=enc)
        except (UnicodeDecodeError, OSError):
            continue
    return ""


def doc_asset_key(rel: str) -> str:
    stem = Path(rel).stem
    parent = "__".join(Path(rel).parent.parts) if Path(rel).parent.parts else "root"
    safe = SAFE_NAME_RE.sub("_", f"{parent}__{stem}").strip("_")
    return safe or "document"


def image_is_useful(path: Path) -> tuple[bool, int, int]:
    try:
        if path.stat().st_size < MIN_IMG_BYTES:
            return False, 0, 0
        from PIL import Image
        with Image.open(path) as im:
            w, h = im.size
            if w < MIN_IMG_DIM or h < MIN_IMG_DIM:
                return False, w, h
            return True, w, h
    except Exception:
        return path.stat().st_size >= MIN_IMG_BYTES, 0, 0


def publish_image(slug: str, doc_key: str, src: Path, name: str, caption: str = "") -> dict | None:
    ok, width, height = image_is_useful(src)
    if not ok:
        return None
    safe_name = SAFE_NAME_RE.sub("_", name).strip("_") or "figure.png"
    if not safe_name.lower().endswith((".png", ".jpg", ".jpeg", ".gif", ".webp")):
        safe_name += ".png"
    rel_dest = f"extracted/{doc_key}/{safe_name}"
    dest = ASSETS_ARCHIVE / slug / rel_dest
    dest.parent.mkdir(parents=True, exist_ok=True)
    try:
        from PIL import Image
        with Image.open(src) as im:
            if im.mode in ("P", "PA", "LA"):
                im = im.convert("RGBA")
            elif im.mode not in ("RGB", "RGBA", "L"):
                im = im.convert("RGB")
            width, height = im.size
            if width < MIN_IMG_DIM or height < MIN_IMG_DIM:
                return None
            if dest.suffix.lower() != ".png" and im.mode == "RGBA":
                dest = dest.with_suffix(".png")
                rel_dest = str(Path(rel_dest).with_suffix(".png"))
            im.save(dest, format="PNG" if dest.suffix.lower() == ".png" else None)
    except Exception:
        if not dest.exists():
            shutil.copy2(src, dest)
    out: dict = {
        "src": f"assets/archive/{slug}/{rel_dest}",
        "alt": caption or safe_name,
        "width": width,
        "height": height,
    }
    if caption:
        out["caption"] = caption
    return out


def merge_segments(segments: list[dict]) -> list[dict]:
    merged: list[dict] = []
    text_buf: list[str] = []
    for seg in segments:
        if seg.get("kind") == "text":
            chunk = (seg.get("content") or "").strip()
            if chunk:
                text_buf.append(chunk)
            continue
        if text_buf:
            merged.append({"kind": "text", "content": "\n\n".join(text_buf)})
            text_buf = []
        merged.append(seg)
    if text_buf:
        deduped: list[str] = []
        prev: str | None = None
        for chunk in text_buf:
            if chunk != prev:
                deduped.append(chunk)
            prev = chunk
        merged.append({"kind": "text", "content": "\n\n".join(deduped)})
    return merged


def cap_segments(segments: list[dict]) -> tuple[list[dict], int]:
    """Limit segments and figure count; return (segments, figure_count)."""
    out: list[dict] = []
    figures = 0
    for seg in segments:
        if len(out) >= MAX_DOC_SEGMENTS:
            break
        if seg.get("kind") == "figures":
            imgs = seg.get("images") or []
            room = MAX_DOC_IMAGES - figures
            if room <= 0:
                continue
            imgs = imgs[:room]
            figures += len(imgs)
            if imgs:
                out.append({"kind": "figures", "images": imgs})
            continue
        out.append(seg)
    return out, figures


def extract_pdf(path: Path) -> str:
    try:
        r = subprocess.run(
            ["pdftotext", "-layout", str(path), "-"],
            capture_output=True, text=True, timeout=60,
        )
        if r.returncode == 0 and r.stdout.strip():
            return r.stdout
    except (OSError, subprocess.TimeoutExpired):
        pass
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        parts = []
        for page in reader.pages:
            parts.append(page.extract_text() or "")
        return "\n\n".join(parts)
    except Exception:
        return ""


def pdf_page_count(path: Path) -> int:
    try:
        r = subprocess.run(["pdfinfo", str(path)], capture_output=True, text=True, timeout=30)
        for line in r.stdout.splitlines():
            if line.startswith("Pages:"):
                return int(line.split(":", 1)[1].strip())
    except (OSError, subprocess.TimeoutExpired, ValueError):
        pass
    return 0


def pdf_images_index(path: Path) -> list[tuple[int, int, int, int]]:
    rows: list[tuple[int, int, int, int]] = []
    try:
        r = subprocess.run(["pdfimages", "-list", str(path)], capture_output=True, text=True, timeout=60)
        for line in r.stdout.splitlines()[2:]:
            if not line.strip():
                continue
            parts = line.split()
            if len(parts) < 5:
                continue
            page, num, typ = int(parts[0]), int(parts[1]), parts[2]
            if typ == "smask":
                continue
            w, h = int(parts[3]), int(parts[4])
            if w >= MIN_IMG_DIM and h >= MIN_IMG_DIM:
                rows.append((page, num, w, h))
    except (OSError, subprocess.TimeoutExpired, ValueError):
        pass
    return rows


def extract_pdf_rich(path: Path, slug: str, doc_key: str) -> tuple[str, list[dict], int]:
    full_text = extract_pdf(path)
    segments: list[dict] = []
    figure_count = 0
    pages = pdf_page_count(path) or 1
    img_rows = pdf_images_index(path)
    imgs_by_page: dict[int, list[tuple[int, int, int]]] = {}
    for page, num, w, h in img_rows:
        imgs_by_page.setdefault(page, []).append((num, w, h))

    try:
        with tempfile.TemporaryDirectory() as td:
            prefix = os.path.join(td, "img")
            subprocess.run(["pdfimages", "-png", str(path), prefix], capture_output=True, timeout=120)
            img_files: dict[int, Path] = {}
            for fname in os.listdir(td):
                m = re.match(r"img-(\d+)\.(\w+)$", fname)
                if m:
                    img_files[int(m.group(1))] = Path(td) / fname

            for page in range(1, pages + 1):
                page_text = ""
                try:
                    r = subprocess.run(
                        ["pdftotext", "-f", str(page), "-l", str(page), "-layout", str(path), "-"],
                        capture_output=True, text=True, timeout=30,
                    )
                    if r.returncode == 0:
                        page_text = clean_text(r.stdout)
                except (OSError, subprocess.TimeoutExpired):
                    pass

                if page_text:
                    segments.append({"kind": "text", "content": page_text})

                page_imgs: list[dict] = []
                for num, w, h in imgs_by_page.get(page, []):
                    src = img_files.get(num)
                    if not src or not src.is_file():
                        continue
                    published = publish_image(
                        slug, doc_key, src, f"page{page:03d}-img{num:03d}.png",
                        caption=f"Page {page}",
                    )
                    if published:
                        page_imgs.append(published)
                if page_imgs:
                    segments.append({"kind": "figures", "images": page_imgs})
                    figure_count += len(page_imgs)
    except (OSError, subprocess.TimeoutExpired):
        pass

    segments = merge_segments(segments)
    segments, figure_count = cap_segments(segments)
    if not full_text.strip() and segments:
        full_text = "\n\n".join(
            s["content"] for s in segments if s.get("kind") == "text" and s.get("content")
        )
    return full_text, segments, figure_count


def extract_doc(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in {".docx", ".dotx"}:
        try:
            from docx import Document
            doc = Document(str(path))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except Exception:
            pass
    try:
        r = subprocess.run(
            ["catdoc", "-w", str(path)],
            capture_output=True, text=True, timeout=60,
        )
        if r.returncode == 0:
            return r.stdout
    except (OSError, subprocess.TimeoutExpired):
        pass
    return ""


def extract_office_html_rich(path: Path, slug: str, doc_key: str) -> tuple[str, list[dict], int]:
    segments: list[dict] = []
    figure_count = 0
    full_text = ""
    try:
        from bs4 import BeautifulSoup
    except ImportError:
        return extract_doc(path), [], 0

    with tempfile.TemporaryDirectory() as td:
        td_path = Path(td)
        staged = td_path / path.name
        shutil.copy2(path, staged)
        r = subprocess.run(
            ["soffice", "--headless", "--convert-to", "html", "--outdir", str(td_path), str(staged)],
            capture_output=True, text=True, timeout=180,
        )
        if r.returncode != 0:
            return extract_doc(path), [], 0

        html_files = list(td_path.glob("*.html"))
        if not html_files:
            return extract_doc(path), [], 0
        html = html_files[0].read_text(encoding="utf-8", errors="replace")
        soup = BeautifulSoup(html, "html.parser")
        body = soup.body or soup
        block_tags = {"p", "h1", "h2", "h3", "h4", "h5", "h6", "li", "td", "th", "pre"}

        for el in body.find_all(["img", *block_tags]):
            if el.name == "img":
                src_name = unquote(el.get("src") or "")
                if not src_name:
                    continue
                src_path = (html_files[0].parent / src_name).resolve()
                if not src_path.is_file():
                    continue
                published = publish_image(slug, doc_key, src_path, Path(src_name).name)
                if published:
                    segments.append({"kind": "figures", "images": [published]})
                    figure_count += 1
                continue

            parents = [p.name for p in el.parents if getattr(p, "name", None) in block_tags]
            if len(parents) > 1:
                continue
            text = clean_text(el.get_text("\n", strip=True))
            if text:
                segments.append({"kind": "text", "content": text})

        full_text = clean_text(body.get_text("\n\n", strip=True))

    segments = merge_segments(segments)
    segments, figure_count = cap_segments(segments)
    if not full_text.strip():
        full_text = extract_doc(path)
    return full_text, segments, figure_count


def extract_ppt(path: Path) -> str:
    try:
        from pptx import Presentation
        prs = Presentation(str(path))
        slides = []
        for i, slide in enumerate(prs.slides, 1):
            bits = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    bits.append(shape.text.strip())
            if bits:
                slides.append(f"Slide {i}\n" + "\n".join(bits))
        return "\n\n".join(slides)
    except Exception:
        return ""


def extract_document_rich(path: Path, slug: str, rel: str) -> tuple[str, list[dict], int]:
    ext = path.suffix.lower()
    name = path.name.lower()
    doc_key = doc_asset_key(rel)

    if ext == ".pdf" or name.endswith(",pdf") or name.endswith(".pdf"):
        return extract_pdf_rich(path, slug, doc_key)
    if ext in OFFICE_HTML_EXT:
        return extract_office_html_rich(path, slug, doc_key)
    if ext in {".ppt", ".pptx"}:
        return extract_ppt(path), [], 0
    if ext == ".rtf":
        raw = read_text_file(path)
        return re.sub(r"\\[a-z]+\d* ?|[{}]", "", raw), [], 0
    return extract_doc(path), [], 0


def section_title(rel: str) -> str:
    parts = Path(rel).parts
    if len(parts) <= 1:
        return "Project root"
    return " / ".join(parts[:-1])


def sort_key(item: dict) -> tuple:
    rel = item["rel"].lower()
    priority = 1
    if "readme" in rel:
        priority = 0
    elif item["type"] == "text" and rel.endswith(".txt"):
        priority = 0
    type_order = {"text": 0, "code": 1, "document": 2, "image": 3, "other": 4}
    return (priority, rel.count("/"), type_order.get(item["type"], 9), rel)


def interleave(items: list[dict]) -> list[dict]:
    texts = [i for i in items if i["type"] == "text"]
    codes = [i for i in items if i["type"] == "code"]
    docs = [i for i in items if i["type"] == "document"]
    images = [i for i in items if i["type"] == "image"]
    others = [i for i in items if i["type"] == "other"]

    readmes = [t for t in texts if "readme" in t["rel"].lower()]
    plain_texts = [t for t in texts if t not in readmes]

    blocks: list[dict] = []
    blocks.extend(readmes)

    queues = [codes, docs, images, plain_texts, others]
    while any(queues):
        for q in queues:
            if q:
                blocks.append(q.pop(0))
    return blocks


def copy_image(slug: str, src: Path, rel: str) -> str:
    dest = ASSETS_ARCHIVE / slug / "images" / rel
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, dest)
    return f"assets/archive/{slug}/images/{rel}"


def process_file(slug: str, company: str, path: Path, rel: str) -> dict | None:
    kind = classify(path)
    if kind == "other":
        return None

    title = Path(rel).name
    block: dict = {
        "type": kind,
        "title": title,
        "rel": rel.replace("\\", "/"),
        "section": section_title(rel),
    }

    if kind == "image":
        block["src"] = copy_image(slug, path, rel.replace("\\", "/"))
        block["alt"] = f"{company} — {title}"
        return block

    if kind == "code":
        raw = read_text_file(path)
        if not raw.strip():
            return None
        content, truncated, total = truncate_lines(raw, MAX_CODE_LINES)
        block["lang"] = LANG_MAP.get(path.suffix.lower(), "text")
        block["content"] = content
        block["truncated"] = truncated
        block["line_count"] = total
        return block

    if kind == "text":
        raw = read_text_file(path)
        if not raw.strip():
            return None
        content, truncated = truncate(clean_text(raw), MAX_TEXT_CHARS)
        block["content"] = content
        block["truncated"] = truncated
        return block

    if kind == "document":
        raw, segments, figure_count = extract_document_rich(path, slug, rel)
        if not raw.strip() and not segments:
            block["content"] = f"[Could not extract text from {title}. File is in the archive.]"
            block["truncated"] = False
            return block
        content, truncated = truncate(clean_text(raw), MAX_DOC_CHARS) if raw.strip() else ("", False)
        block["content"] = content
        block["truncated"] = truncated
        if segments:
            block["segments"] = segments
            block["figure_count"] = figure_count
        return block

    return None


def build_company(company_dir: Path, slug: str) -> dict | None:
    company = company_dir.name
    items: list[dict] = []
    for path in sorted(company_dir.rglob("*")):
        if not path.is_file():
            continue
        rel = str(path.relative_to(company_dir))
        block = process_file(slug, company, path, rel)
        if block:
            items.append(block)

    if not items:
        return None

    ordered = interleave(sorted(items, key=sort_key))
    blocks: list[dict] = []
    current_section: str | None = None
    for item in ordered:
        sec = item.get("section", "Project root")
        if sec != current_section:
            blocks.append({"type": "section", "title": sec})
            current_section = sec
        blocks.append({k: v for k, v in item.items() if k != "section"})

    stats = {
        "code": sum(1 for i in items if i["type"] == "code"),
        "document": sum(1 for i in items if i["type"] == "document"),
        "text": sum(1 for i in items if i["type"] == "text"),
        "image": sum(1 for i in items if i["type"] == "image"),
        "figures": sum(i.get("figure_count", 0) for i in items if i["type"] == "document"),
    }

    return {
        "name": company,
        "slug": slug,
        "stats": stats,
        "blocks": blocks,
    }


def main() -> int:
    if not NNN_ROOT.is_dir():
        print(f"NNN root not found: {NNN_ROOT}", file=sys.stderr)
        MANIFEST.write_text(json.dumps({"employers": {}}, indent=2) + "\n", encoding="utf-8")
        return 1

    employers: dict[str, dict] = {}
    for company_dir in sorted(NNN_ROOT.iterdir()):
        if not company_dir.is_dir():
            continue
        slug = NNN_TO_SLUG.get(company_dir.name)
        if not slug:
            print(f"  skip unmapped folder: {company_dir.name}")
            continue
        entry = build_company(company_dir, slug)
        if entry:
            employers[slug] = entry
            s = entry["stats"]
            print(
                f"  {company_dir.name} → {slug}: "
                f"{s['code']} code, {s['document']} docs, {s['text']} text, "
                f"{s['image']} images, {s['figures']} extracted figures"
            )

    manifest = {
        "source": str(NNN_ROOT),
        "employers": employers,
    }
    MANIFEST.parent.mkdir(parents=True, exist_ok=True)
    MANIFEST.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Wrote {MANIFEST} ({len(employers)} employers)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
